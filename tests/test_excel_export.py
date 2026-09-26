from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.orm import Session


class ExcelExporter:

    def __init__(self, output_dir: str = "data/exports"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(
            parents=True,
            exist_ok=True
        )

    def export(
        self,
        db: Session,
        filename: str = "lead_intelligence.xlsx",
    ) -> Path:

        output_path = self.output_dir / filename

        lead_df = self._get_lead_intelligence(db)
        designer_df = self._get_designer_master(db)
        rug_df = self._get_rug_opportunities(db)

        with pd.ExcelWriter(
            output_path,
            engine="openpyxl"
        ) as writer:

            lead_df.to_excel(
                writer,
                sheet_name="Lead Intelligence Master",
                index=False,
            )

            designer_df.to_excel(
                writer,
                sheet_name="Interior Designer Master",
                index=False,
            )

            rug_df.to_excel(
                writer,
                sheet_name="Rug Opportunity Tracker",
                index=False,
            )

            self._format_workbook(writer)

        return output_path

    def _format_workbook(self, writer):

        workbook = writer.book

        for worksheet in workbook.worksheets:

            # Freeze header row
            worksheet.freeze_panes = "A2"

            # Enable filters
            if worksheet.max_row >= 1:
                worksheet.auto_filter.ref = worksheet.dimensions

            # Header formatting
            for cell in worksheet[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

            # Body formatting
            for row in worksheet.iter_rows(
                min_row=2
            ):
                for cell in row:
                    cell.alignment = Alignment(
                        vertical="top",
                        wrap_text=True,
                    )

            # Automatic column width
            for column_cells in worksheet.columns:

                max_length = 0
                column_letter = get_column_letter(
                    column_cells[0].column
                )

                for cell in column_cells:
                    if cell.value is not None:
                        max_length = max(
                            max_length,
                            len(str(cell.value))
                        )

                # Keep widths reasonable
                worksheet.column_dimensions[
                    column_letter
                ].width = min(
                    max(max_length + 2, 12),
                    45
                )

            # Header row height
            worksheet.row_dimensions[1].height = 30

    def _get_lead_intelligence(
        self,
        db: Session
    ) -> pd.DataFrame:

        query = text("""
            SELECT
                p.id AS project_id,
                a.url AS article_url,
                a.title AS article_title,
                a.published_at,
                p.project_name,
                p.home_type,
                p.location,
                p.project_size,
                p.completion_year,
                p.homeowner_name,
                p.homeowner_profession,
                p.homeowner_industry,
                p.designer_name,
                p.designer_studio,
                p.designer_role,
                p.designer_city,
                p.designer_website,
                p.flooring,
                p.furniture,
                p.decor,
                p.textile_elements,
                p.sourcing_notes,
                s.lead_score,
                s.rug_opportunity_score
            FROM projects p
            JOIN articles a
                ON p.article_id = a.id
            LEFT JOIN scores s
                ON p.id = s.project_id
            ORDER BY a.published_at DESC NULLS LAST
        """)

        result = db.execute(query)

        return pd.DataFrame(
            result.mappings().all()
        )

    def _get_designer_master(
        self,
        db: Session
    ) -> pd.DataFrame:

        query = text("""
            SELECT
                id,
                designer_name,
                studio_name,
                website,
                address,
                city,
                country,
                contact,
                email,
                company_description,
                linkedin_url,
                instagram_url,
                facebook_url,
                youtube_url,
                last_enriched_at
            FROM designers
            ORDER BY designer_name
        """)

        result = db.execute(query)

        return pd.DataFrame(
            result.mappings().all()
        )

    def _get_rug_opportunities(
        self,
        db: Session
    ) -> pd.DataFrame:

        query = text("""
            SELECT
                r.id AS rug_id,
                r.project_id,
                p.project_name,
                p.designer_name,
                p.designer_studio,
                p.location,
                r.rug_used,
                r.rug_type,
                r.origin,
                r.material,
                r.supplier,
                r.brand,
                r.handmade,
                r.handwoven,
                r.vintage,
                r.custom_made,
                r.imported,
                r.sourcing_notes,
                r.opportunity_score
            FROM rugs r
            JOIN projects p
                ON r.project_id = p.id
            ORDER BY r.opportunity_score DESC NULLS LAST
        """)

        result = db.execute(query)

        return pd.DataFrame(
            result.mappings().all()
        )